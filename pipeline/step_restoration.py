"""Step-restoration analysis (standalone; read-only; does not touch the frozen app).

Follows the AEP HackOHI/O data guide:
  §3  ARCOS damage -> nearest pole (height class) -> workbook labor -> x4.5 = PERSON-hours
  §4  Step restoration: isolate damaged sections with fuses/reclosers/switches; who depends on what;
      which sections can be restored first; critical facilities influence priority
  §5  Crews of 2-8 people; compare restoring everyone in 2 days / 3 days / 1 week;
      restoration effort (person-hours, crew size, dependencies) vs value (customers, criticality)
  Key questions: 1 What is damaged? 2 Who depends on it? 3 What work? 4 What first, how long?

Part A: Demo location A (the app's simulated fuse + transformer damage).
Part B: A real ARCOS storm on Demo A's circuit (72), circuit-wide.

Every number comes from the AEP files; every choice the data doesn't make is listed as an ASSUMPTION.
Privacy: prints counts, hours and distances only; never premise IDs, addresses or facility names.

Run from repo root:  python pipeline/step_restoration.py > step_output.txt
"""
import heapq
import re
import statistics

import geopandas as gpd
import openpyxl
import pandas as pd

D = "data/extracted/"
WORKBOOK = "data/raw/Install_Labor_HackOHIO.xlsx"
CIRCUIT_TOKEN_ID = "157964494"      # demo-a's circuit breaker (from backend.data.prepare_context)
CIRCUIT_CODE = "0000072"
DEMO_TRANSFORMER = "142761938"
DEMO_FUSE = "137975333"

CREW_SIZES = [2, 4, 8]              # guide §5: "a reasonable crew may contain 2-8 people"
CREW_COUNTS = [1, 2, 3, 4, 6, 8]
TARGETS_H = {"2 days": 48, "3 days": 72, "1 week": 168}
ASSUMPTIONS = {
    "around_the_clock": "Crews work in rotating shifts, so elapsed hours run continuously from when crews start.",
    "nearest_premise": "Each damage point affects the section of its nearest circuit-72 customer (guide §3: nearest spatial match).",
    "isolation": "Transformer/secondary damage is isolated at that transformer; primary/pole damage at the nearest upstream "
                 "fuse/recloser/switch; damage to an isolating device itself at the next device upstream; none -> whole circuit.",
    "capacitor": "Capacitor bank damage needs repair work but does not interrupt customers.",
    "no_rate": "Primary and secondary conductor repairs have no workbook rate: 2.0 and 1.0 direct hours (team assumption).",
    "pole": "Broken pole = remove + replace pole and crossarm (guide §3); size from nearest pole height within 30 m, else Medium.",
    "transformer": "Damaged transformer size unknown: 100 kVA workbook row used (same as the app).",
    "fuse_bank": "One cutout per damaged fuse bank (phase count not in ARCOS record).",
    "access": "ARCOS Access_Iss = Yes adds the workbook's 0.5 h restricted-access labor once per damage record.",
    "critical": "A critical facility is served like its nearest customer if that customer is within 150 m (guide §4: "
                "'create reasonable spatial relationships'). Each counts as 100 customers when prioritizing.",
    "priority": "Repair order = upstream first (breaker, then reclosers, then fuses, then transformers), and within each "
                "level the most (customers + critical weight) per person-hour. A customer is restored only when every "
                "damaged section on its path is repaired. Each damage item is a separate job, so several crews can work on one "
                "section at the same time.",
    "upper": "Time ranges use the app's v2 envelope: lower = person-hours / crew size (everyone fully productive), upper = person-hours / 1 (one effective worker). Real efficiency is unknown.",
}
CRITICAL_RADIUS_M, CRITICAL_WEIGHT, POLE_MATCH_M = 150, 100, 30


# ---------------- labor from the AEP workbook (guide §3) ----------------
def read_labor(path=WORKBOOK):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = lambda name: next(wb[s] for s in wb.sheetnames if s.strip().lower() == name.lower())
    mat, only, ilm = sheet("Labor time material related"), sheet("Labor only work"), sheet("Indirect Labor Multiplier")
    num = lambda ws, c: ws[c].value if isinstance(ws[c].value, (int, float)) else None
    pair = lambda ws, a, b: (num(ws, a), num(ws, b))
    L = {"ilm": num(ilm, "B1"), "access": num(only, "D2"),
         "crossarm": pair(mat, "A3", "B3"), "fuse_cutout": pair(mat, "A90", "B90"),
         "transformer_100": pair(mat, "C25", "D25"), "recloser": pair(mat, "A70", "B70"),
         "switch": pair(mat, "A85", "B85"),
         "pole": {"Small": pair(mat, "C12", "D12"), "Medium": pair(mat, "C13", "D13"), "Large": pair(mat, "C14", "D14")}}
    return L


def direct_hours(L, asset, pole_size="Medium"):
    """Direct (hands-on) install+remove hours for one damaged item; returns (hours, source)."""
    tot = lambda p: None if None in p else p[0] + p[1]
    a = (asset or "").strip().lower()
    if a == "pole":
        h = tot(L["pole"][pole_size]) + (tot(L["crossarm"]) or 0)
        return h, f"workbook pole ({pole_size}) + crossarm"
    if a == "transformer":
        return tot(L["transformer_100"]), "workbook 100 kVA transformer"
    if a == "fuse bank":
        return tot(L["fuse_cutout"]), "workbook fuse cutout"
    if a == "recloser bank" and tot(L["recloser"]) is not None:
        return tot(L["recloser"]), "workbook recloser"
    if a == "switch" and tot(L["switch"]) is not None:
        return tot(L["switch"]), "workbook switch"
    if a == "primary":
        return 2.0, "ASSUMPTION (no workbook rate)"
    if a == "secondary":
        return 1.0, "ASSUMPTION (no workbook rate)"
    return 2.0, "ASSUMPTION (no workbook rate)"


def pole_size(height):
    try:
        h = float(height)
    except (TypeError, ValueError):
        return "Medium"
    return "Small" if h <= 37 else "Medium" if h < 58 else "Large"


# ---------------- network topology (guide §2) ----------------
def load_chains(path=D + "matched_upstream/matched_upstream.txt"):
    """premise -> ordered list of (kind, id), closest equipment first."""
    chains = {}
    with open(path) as f:
        for line in f:
            parts = line.rstrip("\n").split("\t")
            toks = [p.split("*") for p in parts[1:] if p.count("*") == 3]
            if toks:
                chains[parts[0].strip().zfill(9)] = [(t[2], t[1]) for t in toks]
    return chains


def is_isolating(kind):
    return bool(re.search(r"fuse|recloser|switch|sectionaliz", kind, re.I))


def behind_index(chains):
    idx = {}
    for p, chain in chains.items():
        for _, dev in chain:
            idx.setdefault(dev, set()).add(p)
    return idx


def section_for(chain, asset):
    """Which device isolates this damage? Returns (device_id, kind)."""
    a = (asset or "").strip().lower()
    if a in ("transformer", "secondary"):
        return chain[0][1], chain[0][0]
    iso = [(k, d) for k, d in chain if is_isolating(k)]
    skip = 1 if a in ("fuse bank", "recloser bank", "switch") else 0
    if len(iso) > skip:
        return iso[skip][1], iso[skip][0]
    return CIRCUIT_TOKEN_ID, "Circuit Breaker (whole circuit)"


# ---------------- scheduling (guide §4-5) ----------------
def device_depth(chains):
    """How many protective devices sit between each device and the circuit breaker (breaker = 0)."""
    depth = {CIRCUIT_TOKEN_ID: 0}
    for chain in chains.values():
        upstream = 0
        for kind, dev in reversed(chain[:-1]):          # walk from the breaker end toward the customer
            if is_isolating(kind):
                upstream += 1
            d = upstream if is_isolating(kind) else upstream + 1   # transformer sits below its fuse
            depth[dev] = min(depth.get(dev, d), d)
    return depth


def schedule(sections, cust_sections, crews, crew_size, critical_by_section, depth):
    """List-schedule sections on `crews` crews of `crew_size` people. Returns per-customer restore hour.
    Upstream sections first (guide §4 step restoration); within a level, most value per person-hour."""
    order = sorted(sections, key=lambda s: (depth.get(s, 99) if sections[s]["impacts"] else 999,
                   -(sections[s]["dependents"] + CRITICAL_WEIGHT * critical_by_section.get(s, 0))
                   / max(sections[s]["person_hours"], 1e-9)))
    free = [0.0] * crews
    heapq.heapify(free)
    done = {}
    for s in order:                      # each damage item is its own job; a section is done when all its items are
        finish = 0.0
        for job in sorted(sections[s]["jobs"], reverse=True):
            start = heapq.heappop(free)
            end = start + job / crew_size
            finish = max(finish, end)
            heapq.heappush(free, end)
        done[s] = finish
    restore = {p: max(done[s] for s in secs) for p, secs in cust_sections.items()}
    return order, done, restore


def pct_by(restore, hours):
    vals = list(restore.values())
    return 100.0 * sum(v <= hours for v in vals) / len(vals) if vals else 100.0


# ---------------- main analysis ----------------
def main():
    L = read_labor()
    print("== LABOR READ FROM WORKBOOK (guide §3)")
    print(f"ILM (Indirect Labor Multiplier!B1): {L['ilm']} | restricted access (Labor only work!D2): {L['access']}")
    print(f"fuse cutout A90:B90 {L['fuse_cutout']} | 100 kVA transformer C25:D25 {L['transformer_100']} | "
          f"crossarm A3:B3 {L['crossarm']} | recloser A70:B70 {L['recloser']} | switch A85:B85 {L['switch']}")
    print(f"poles C12:D14 {L['pole']}")
    ILM, ACCESS = L["ilm"] or 4.5, L["access"] or 0.5

    chains = load_chains()
    idx = behind_index(chains)
    circuit = idx.get(CIRCUIT_TOKEN_ID, set())
    chains72 = {p: chains[p] for p in circuit}
    depth = device_depth(chains72)

    # -------- Part A: Demo location A --------
    on_fuse, on_x = idx.get(DEMO_FUSE, set()), idx.get(DEMO_TRANSFORMER, set())
    print("\n================ PART A: DEMO LOCATION A (app's simulated fuse + transformer damage) ================")
    print(f"Q2 Who depends on it?  Circuit 72: {len(circuit)} customers | behind damaged fuse: {len(on_fuse)} | "
          f"behind damaged transformer: {len(on_x)} (incl. Demo A)")
    print(f"Can it be isolated?    The fuse is a potential isolation boundary for its {len(on_fuse)}-customer branch; "
          f"the other {len(circuit - on_fuse)} circuit-72 customers are outside the damaged branch "
          f"(assumes the fuse operates; switching state is not in the data).")
    print(f"Restore first?         Fix fuse first -> {len(on_fuse - on_x)} customer(s) back; then transformer -> "
          f"{len(on_x)} (Demo A).")
    fuse_ph = direct_hours(L, "fuse bank")[0] * ILM
    xf_ph = direct_hours(L, "transformer")[0] * ILM
    for size in CREW_SIZES:
        t1, t2 = fuse_ph / size, (fuse_ph + xf_ph) / size
        print(f"  Crew of {size}: fuse done {t1:.1f}-{t1*size:.1f} h | Demo A restored {t2:.1f}-{t2*size:.1f} h "
              f"(should match the app v2 envelope)")

    # -------- Part B: real ARCOS storm on circuit 72 --------
    cust = gpd.read_file(D + "Data1/Customer_Clip.shp")[["premise_nu", "geometry"]]
    cust["premise"] = cust["premise_nu"].astype(str).str.strip().str.replace(r"\.0$", "", regex=True).str.zfill(9)
    cust = cust[cust["premise"].isin(chains72)].drop_duplicates("premise")
    utm = cust.estimate_utm_crs()
    cust = cust.to_crs(utm)[["premise", "geometry"]]

    dmg = gpd.read_file(D + "Data1/arcos_prod_damage_09_28_Clip.shp").to_crs(utm)
    dmg["circ"] = dmg["Circuit"].astype(str).str.extract(r"(\d+)")[0].fillna("").str.zfill(7)
    dmg["report"] = pd.to_datetime(dmg["Report_Tim"], errors="coerce")
    dmg["repaired"] = pd.to_datetime(dmg["Repaired_D"], errors="coerce")
    dmg["storm"] = dmg["report"].apply(lambda t: "Unknown" if pd.isna(t) else "Sep 2024" if t < pd.Timestamp("2025-01-01") else "Feb 2025")
    dmg = dmg[dmg["circ"] == CIRCUIT_CODE].reset_index(drop=True)

    poles = gpd.read_file(D + "Data1/Pole_Clip.shp")[["height", "geometry"]].to_crs(utm)
    fac = []
    for label, fname in [("Hospital", "Hospitals_Medic_Pairwis_Clip.shp"), ("Fire station", "Fire_Stations_E_Pairwis_Clip.shp"),
                         ("Ambulance", "Ambulance_Servi_Pairwis_Clip.shp")]:
        g = gpd.read_file(D + "Data1/" + fname)[["geometry"]].to_crs(utm)
        g["type"] = label
        fac.append(g)
    fac = gpd.GeoDataFrame(pd.concat(fac, ignore_index=True), crs=utm)
    fac = gpd.sjoin_nearest(fac, cust, how="left", distance_col="dist_m").drop_duplicates(subset=["geometry"])
    fac72 = fac[fac["dist_m"] <= CRITICAL_RADIUS_M]

    print("\n================ PART B: REAL ARCOS STORM DAMAGE ON CIRCUIT 72 (circuit-wide) ================")
    print("Damage records on circuit 72 by storm:", {k: int(v) for k, v in dmg["storm"].value_counts().items()})
    print(f"Critical facilities tied to circuit-72 customers (within {CRITICAL_RADIUS_M} m): "
          f"{ {k: int(v) for k, v in fac72['type'].value_counts().items()} or 'none'}")

    for storm in [s for s in ("Feb 2025", "Sep 2024") if (dmg["storm"] == s).any()]:
        d = dmg[dmg["storm"] == storm].copy()
        d = gpd.sjoin_nearest(d, cust, how="left", distance_col="cust_dist").drop_duplicates(subset=d.columns.drop("geometry").tolist() + ["geometry"])
        is_pole = d["Asset_Type"].astype(str).str.strip().str.lower() == "pole"
        sizes = pd.Series("Medium", index=d.index)
        if is_pole.any():
            pj = gpd.sjoin_nearest(d.loc[is_pole, ["geometry"]], poles, how="left", max_distance=POLE_MATCH_M)
            pj = pj[~pj.index.duplicated()]
            sizes.loc[pj.index] = pj["height"].apply(pole_size)

        sections, damage_rows = {}, []
        for i, r in d.iterrows():
            if pd.isna(r.get("premise")):
                continue
            asset = str(r["Asset_Type"])
            h, src = direct_hours(L, asset, sizes.get(i, "Medium"))
            h += ACCESS if str(r.get("Access_Iss")) == "Yes" else 0
            dev, kind = section_for(chains72[r["premise"]], asset)
            impacts = asset.strip().lower() != "capacitor bank"
            if not impacts:
                dev, kind = f"cap-{i}", "Capacitor bank (no outage)"
            s = sections.setdefault(dev, {"kind": kind, "items": 0, "person_hours": 0.0, "impacts": False, "jobs": []})
            s["jobs"].append(h * ILM)
            s["items"] += 1
            s["person_hours"] += h * ILM
            s["impacts"] |= impacts
            damage_rows.append((asset, src, round(r["cust_dist"])))

        impacting = {sid for sid, s in sections.items() if s["impacts"]}
        cust_sections = {}
        for p, chain in chains72.items():
            ids = {dev for _, dev in chain} | {CIRCUIT_TOKEN_ID}
            secs = ids & impacting
            if secs:
                cust_sections[p] = secs
        for sid, s in sections.items():
            s["dependents"] = sum(1 for secs in cust_sections.values() if sid in secs)
        crit_by_section = {}
        for _, f in fac72.iterrows():
            for sid in cust_sections.get(f["premise"], ()):
                crit_by_section[sid] = crit_by_section.get(sid, 0) + 1
        crit_customers = {f["premise"]: f["type"] for _, f in fac72.iterrows() if f["premise"] in cust_sections}

        total_ph = sum(s["person_hours"] for s in sections.values())
        hist = (d["repaired"] - d["report"]).dt.total_seconds() / 3600
        hist = hist[(hist >= 0) & (hist <= 30 * 24)]
        print(f"\n---------- Storm: {storm} ----------")
        print(f"Q1 What is damaged?  {len(d)} records: { {k: int(v) for k, v in d['Asset_Type'].value_counts().items()} }; "
              f"median distance to nearest customer {statistics.median([x[2] for x in damage_rows]) if damage_rows else 'n/a'} m")
        print(f"   Unrated items using team assumptions: {sum('ASSUMPTION' in x[1] for x in damage_rows)} of {len(damage_rows)}")
        print(f"Q2 Who depends on it?  {len(cust_sections)} of {len(chains72)} circuit-72 customers lose power "
              f"({100*len(cust_sections)/max(len(chains72),1):.0f}%); {len(chains72)-len(cust_sections)} keep power")
        print(f"   Critical facilities affected: { {k: int(v) for k, v in pd.Series(list(crit_customers.values()), dtype=object).value_counts().items()} or 'none'}")
        print(f"Can it be isolated?  Damage falls into {len(impacting)} isolatable sections "
              f"({sum(1 for s in impacting if s == CIRCUIT_TOKEN_ID)} require the whole circuit):")
        for sid in sorted(impacting, key=lambda x: -sections[x]["dependents"])[:8]:
            s = sections[sid]
            print(f"   - {s['kind']:<32} {s['items']:>3} damage items, {s['person_hours']:6.1f} person-h, "
                  f"{s['dependents']:>4} customers depend on it, critical: {crit_by_section.get(sid, 0)}")
        print(f"Q3 What work?  {total_ph:.0f} total person-hours (direct x {ILM}); "
              f"historical median repair time for these records: {hist.median():.0f} h" if len(hist) else "")

        print(f"Q4 What first, and how long?  (4 crews x 4 people)")
        order, done, restore = schedule(sections, cust_sections, 4, 4, crit_by_section, depth)
        restored_so_far = set()
        by_finish = sorted(order, key=lambda x: done[x])
        print(f"   Priority (start) order: " + ", ".join(sections[x]["kind"].split(" (")[0] for x in order[:8]))
        for n, sid in enumerate(by_finish[:10], 1):
            newly = {p for p, t in restore.items() if t <= done[sid]} - restored_so_far
            restored_so_far |= newly
            crit = "  <- critical facility" if any(p in crit_customers for p in newly) else ""
            print(f"   Finish {n:>2}: {sections[sid]['kind']:<28} done {done[sid]:5.1f}-{done[sid]*4:5.1f} h "
                  f"-> {len(newly):>4} customers restored (total {len(restored_so_far)}){crit}")
        if len(by_finish) > 10:
            print(f"   ... {len(order)-10} more sections")
        for p, ftype in crit_customers.items():
            print(f"   {ftype} restored at {restore[p]:.1f}-{restore[p]*4:.1f} h")

        print(f"Scenarios (guide §5): time until ALL affected customers are restored, hours (lower = full crew, upper = one effective worker)")
        print("   crews  " + "  ".join(f"{s}-person crews".rjust(18) for s in CREW_SIZES))
        need = {k: {} for k in TARGETS_H}
        for c in CREW_COUNTS:
            cells = []
            for size in CREW_SIZES:
                _, _, rest = schedule(sections, cust_sections, c, size, crit_by_section, depth)
                t = max(rest.values()) if rest else 0
                cells.append(f"{t:6.1f}-{t*size:6.1f} h".rjust(18))
                for k, lim in TARGETS_H.items():
                    if t <= lim and size not in need[k]:
                        need[k][size] = c
            print(f"   {c:>5}  " + "  ".join(cells))
        for k in TARGETS_H:
            parts = [f"{need[k].get(s, f'>{CREW_COUNTS[-1]}')} crews of {s}" for s in CREW_SIZES]
            print(f"   Restore everyone within {k} (full-crew productivity): " + " | ".join(parts))
        print(f"   With 4 crews of 4: {pct_by(restore, 24):.0f}% restored by 24 h, {pct_by(restore, 48):.0f}% by 2 days, "
              f"{pct_by(restore, 72):.0f}% by 3 days, {pct_by(restore, 168):.0f}% by 1 week")

    print("\n== ASSUMPTIONS (state these to judges)")
    for k, v in ASSUMPTIONS.items():
        print(f"- {v}")
    print("- Damage is a historical ARCOS storm replayed as a scenario; this is not live outage or official AEP ETR information.")


if __name__ == "__main__":
    main()